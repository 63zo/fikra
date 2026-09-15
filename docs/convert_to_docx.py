import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# University of Bisha Colors
COLOR_PRIMARY = RGBColor(20, 87, 58)    # #14573A Dark Green
COLOR_GOLD = RGBColor(197, 160, 89)     # #C5A059 Gold
COLOR_DARK = RGBColor(30, 41, 59)       # #1E293B Slate
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_PRIMARY = "14573A"
HEX_GOLD = "C5A059"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "E2E8F0"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_rtl(paragraph_or_run):
    pPr = paragraph_or_run._element.get_or_add_pPr() if hasattr(paragraph_or_run, '_element') else paragraph_or_run._p.get_or_add_pPr()
    bidi = parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>')
    pPr.append(bidi)

def is_arabic(text):
    return any(ord(char) >= 0x0600 and ord(char) <= 0x06FF for char in text)

def apply_formatting(run, font_name="Calibri", font_size=11, bold=False, italic=False, color=COLOR_DARK):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    # Set complex script font for Arabic
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="Traditional Arabic"/>')
    rPr.append(rFonts)

def parse_inline_markdown(paragraph, text, base_font="Calibri", base_size=11, default_color=COLOR_DARK):
    # Regex for bold **text**, italic *text*, inline `code`, [link](url)
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|[^*`\[]+)'
    tokens = re.findall(pattern, text)
    if not tokens:
        tokens = [text]

    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            apply_formatting(run, font_name=base_font, font_size=base_size, bold=True, color=default_color)
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            apply_formatting(run, font_name=base_font, font_size=base_size, italic=True, color=default_color)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            apply_formatting(run, font_name="Consolas", font_size=base_size - 1, bold=True, color=COLOR_PRIMARY)
        elif token.startswith('[') and '](' in token:
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                label, url = m.groups()
                run = paragraph.add_run(label)
                apply_formatting(run, font_name=base_font, font_size=base_size, bold=True, color=COLOR_PRIMARY)
            else:
                run = paragraph.add_run(token)
                apply_formatting(run, font_name=base_font, font_size=base_size, color=default_color)
        else:
            run = paragraph.add_run(token)
            apply_formatting(run, font_name=base_font, font_size=base_size, color=default_color)

def convert_md_to_docx(md_path, docx_path):
    print(f"Converting: {os.path.basename(md_path)} -> {os.path.basename(docx_path)}")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Document Title Header
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("جامعة بيشة | University of Bisha — منصة فكرة (Fikra)")
    apply_formatting(hrun, font_name="Calibri", font_size=9, color=COLOR_MUTED)

    # Document Footer
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("منصة إدارة الابتكار والأفكار المؤسسية © 2026")
    apply_formatting(frun, font_name="Calibri", font_size=9, color=COLOR_MUTED)

    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""

    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Handle Code blocks & Mermaid diagrams
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                in_code_block = False
                # Render code block box
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                table.columns[0].width = Inches(6.5)
                
                cell = table.cell(0, 0)
                set_cell_background(cell, "F1F5F9" if code_lang != "mermaid" else "F0FDF4")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                # Title of block
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(4)
                lbl = f"[{code_lang.upper() if code_lang else 'CODE'}]"
                lbl_run = cp.add_run(lbl + "\n")
                apply_formatting(lbl_run, font_name="Consolas", font_size=9, bold=True, color=COLOR_PRIMARY if code_lang == "mermaid" else COLOR_MUTED)

                code_text = "\n".join(code_lines)
                code_run = cp.add_run(code_text)
                apply_formatting(code_run, font_name="Consolas", font_size=9.5, color=COLOR_DARK)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith('|') and '|' in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table
            raw_rows = []
            for tline in table_lines:
                # ignore separator row like |---|---|
                cells = [c.strip() for c in tline.split('|')[1:-1]]
                if cells and all(re.match(r'^[:\s\-]+$', c) for c in cells):
                    continue
                raw_rows.append(cells)

            if raw_rows:
                num_cols = max(len(r) for r in raw_rows)
                tbl = doc.add_table(rows=len(raw_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl)

                for r_idx, row in enumerate(raw_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell_val = row[c_idx] if c_idx < len(row) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        
                        if is_header:
                            set_cell_background(cell, HEX_PRIMARY)
                        elif r_idx % 2 == 1:
                            set_cell_background(cell, HEX_LIGHT_BG)
                        else:
                            set_cell_background(cell, "FFFFFF")

                        cp = cell.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(3)
                        cp.paragraph_format.space_after = Pt(3)
                        if is_arabic(cell_val):
                            cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            set_rtl(cp)

                        parse_inline_markdown(
                            cp, cell_val,
                            base_font="Calibri",
                            base_size=10,
                            default_color=COLOR_WHITE if is_header else COLOR_DARK
                        )
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Handle Alerts (> [!NOTE], > [!CAUTION], etc.)
        if stripped.startswith('>'):
            alert_text = stripped.lstrip('> ').strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "FEF3C7" if "CAUTION" in alert_text or "WARNING" in alert_text else "EFF6FF")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            
            cp = cell.paragraphs[0]
            if is_arabic(alert_text):
                cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(cp)
            parse_inline_markdown(cp, alert_text, base_font="Calibri", base_size=10.5, default_color=RGBColor(146, 64, 14) if "CAUTION" in alert_text else RGBColor(30, 58, 138))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Handle Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            if is_arabic(stripped):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)
            parse_inline_markdown(p, stripped[2:], base_font="Arial", base_size=20, default_color=COLOR_PRIMARY)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            if is_arabic(stripped):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)
            parse_inline_markdown(p, stripped[3:], base_font="Arial", base_size=15, default_color=COLOR_PRIMARY)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            if is_arabic(stripped):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)
            parse_inline_markdown(p, stripped[4:], base_font="Arial", base_size=12.5, default_color=COLOR_GOLD)
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            if is_arabic(stripped):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)
            parse_inline_markdown(p, stripped[5:], base_font="Calibri", base_size=11.5, default_color=COLOR_DARK)
        elif stripped.startswith('---'):
            # Horizontal divider line
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("―" * 45)
            apply_formatting(r, font_name="Calibri", font_size=10, color=COLOR_GOLD)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            if is_arabic(stripped):
                set_rtl(p)
            parse_inline_markdown(p, stripped[2:], base_font="Calibri", base_size=11, default_color=COLOR_DARK)
        elif re.match(r'^\d+\.\s+', stripped):
            m = re.match(r'^(\d+\.\s+)(.*)$', stripped)
            num_prefix, rest = m.groups()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            if is_arabic(stripped):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)
            r_num = p.add_run(num_prefix)
            apply_formatting(r_num, font_name="Calibri", font_size=11, bold=True, color=COLOR_PRIMARY)
            parse_inline_markdown(p, rest, base_font="Calibri", base_size=11, default_color=COLOR_DARK)
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            if is_arabic(stripped):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)
            parse_inline_markdown(p, stripped, base_font="Calibri", base_size=11, default_color=COLOR_DARK)

        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")

def main():
    docs_dir = r"c:\Users\ezzoa\OneDrive\Documents\Desktop\Projects\bsha\docs"
    md_files = [
        ("01_USER_MANUAL.md", "01_USER_MANUAL.docx"),
        ("02_SYSTEM_ADMINISTRATION_GUIDE.md", "02_SYSTEM_ADMINISTRATION_GUIDE.docx"),
        ("03_SOFTWARE_ARCHITECTURE_DOCUMENT.md", "03_SOFTWARE_ARCHITECTURE_DOCUMENT.docx"),
        ("04_API_AND_INTEGRATION_DOCUMENTATION.md", "04_API_AND_INTEGRATION_DOCUMENTATION.docx"),
        ("05_DATABASE_DESIGN_DOCUMENT.md", "05_DATABASE_DESIGN_DOCUMENT.docx"),
        ("README.md", "00_FIKRA_DOCUMENTATION_HUB.docx"),
    ]

    for md_name, docx_name in md_files:
        md_full = os.path.join(docs_dir, md_name)
        docx_full = os.path.join(docs_dir, docx_name)
        if os.path.exists(md_full):
            convert_md_to_docx(md_full, docx_full)

    print("\nAll DOCX files generated successfully!")

if __name__ == "__main__":
    main()
